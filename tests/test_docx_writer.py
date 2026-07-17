from zipfile import ZipFile

from docx import Document

from device_report.docx_writer import write_docx
from device_report.models import Diagnostic, Field, ReportData, Section, Table


def sample_report(language="en"):
    return ReportData(
        language=language,
        title="Lenovo ThinkPad T14",
        generated_at="2026-07-17 18:30:00",
        sections=[
            Section(
                key="overview",
                title="Device overview" if language == "en" else "Обзор устройства",
                fields=[Field("Manufacturer" if language == "en" else "Производитель", "Lenovo")],
            ),
            Section(
                key="cpu",
                title="Processor" if language == "en" else "Процессор",
                tables=[Table(headers=("Name", "Cores"), rows=(("Ryzen", 8),))],
            ),
            Section(key="audio", title="Audio devices" if language == "en" else "Аудиоустройства"),
        ],
        diagnostics=[Diagnostic(source="battery", message="CommandTimeout")],
    )


def document_text(document):
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def test_writer_creates_structured_english_document(tmp_path):
    output = tmp_path / "report.docx"

    result = write_docx(sample_report("en"), output)

    assert result == output
    document = Document(output)
    text = document_text(document)
    assert "Lenovo ThinkPad T14" in text
    assert "Report date: 2026-07-17 18:30:00" in text
    assert "Device overview" in text
    assert "Processor" in text
    assert "Audio devices" in text
    assert "Unavailable" in text
    assert "Diagnostics" in text
    assert "CommandTimeout" in text
    assert "Ryzen" in text


def test_writer_localizes_document_labels(tmp_path):
    output = tmp_path / "report-ru.docx"

    write_docx(sample_report("ru"), output)

    text = document_text(Document(output))
    assert "Дата отчёта: 2026-07-17 18:30:00" in text
    assert "Недоступно" in text
    assert "Диагностика" in text


def test_writer_adds_toc_page_fields_headers_and_repeated_table_headers(tmp_path):
    output = tmp_path / "report.docx"
    write_docx(sample_report(), output)

    with ZipFile(output) as archive:
        xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )

    assert 'TOC \\o "1-3"' in xml
    assert "PAGE" in xml
    assert "tblHeader" in xml
    assert "D9EAF7" in xml


def test_writer_does_not_leak_excluded_sentinel_text(tmp_path):
    report = sample_report()
    report.sections[0].fields.append(Field("Description", "SAFE"))
    report.sections[0].fields.append(Field("SerialNumber", "SUPER_SECRET_SENTINEL"))
    output = tmp_path / "report.docx"

    write_docx(report, output)

    with ZipFile(output) as archive:
        package = b"".join(archive.read(name) for name in archive.namelist())
    assert b"SUPER_SECRET_SENTINEL" not in package


def test_writer_does_not_list_enabled_tabs_on_cover(tmp_path):
    report = sample_report()
    report.selected_sections = ("ENABLED_TAB_SENTINEL",)
    output = tmp_path / "report.docx"

    write_docx(report, output)

    with ZipFile(output) as archive:
        package = b"".join(archive.read(name) for name in archive.namelist())
    assert b"ENABLED_TAB_SENTINEL" not in package
