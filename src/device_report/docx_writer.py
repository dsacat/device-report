from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .i18n import translate
from .models import ReportData, Section, Table
from .privacy import sanitize_mapping


ACCENT = "1F4E78"
LIGHT_ACCENT = "D9EAF7"
LIGHT_ROW = "F5F8FA"


def _set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _repeat_table_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _add_field(paragraph: Any, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction_node, separate, end))


def _configure_styles(document: DocumentType) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(4)
    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(ACCENT)
    if "Report Subtitle" not in styles:
        subtitle = styles.add_style("Report Subtitle", WD_STYLE_TYPE.PARAGRAPH)
        subtitle.font.name = "Aptos"
        subtitle.font.size = Pt(11)
        subtitle.font.color.rgb = RGBColor(90, 100, 110)
        subtitle.paragraph_format.space_after = Pt(12)


def _configure_sections(document: DocumentType, report: ReportData) -> None:
    for section in document.sections:
        section.top_margin = Cm(1.7)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)
        header = section.header.paragraphs[0]
        header.text = report.title
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if header.runs:
            header.runs[0].font.size = Pt(8)
            header.runs[0].font.color.rgb = RGBColor(110, 120, 130)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run(f"{report.generated_at}  •  ")
        _add_field(footer, "PAGE")


def _safe_field_items(section: Section) -> list[tuple[str, Any]]:
    safe: list[tuple[str, Any]] = []
    for item in section.nonempty_fields():
        sanitized = sanitize_mapping({item.key: item.value})
        if item.key in sanitized:
            safe.append((item.key, sanitized[item.key]))
    return safe


def _safe_table(table: Table) -> tuple[tuple[str, ...], tuple[tuple[Any, ...], ...]]:
    kept: list[int] = []
    for index, header in enumerate(table.headers):
        if header in sanitize_mapping({header: "probe"}):
            kept.append(index)
    headers = tuple(table.headers[index] for index in kept)
    rows: list[tuple[Any, ...]] = []
    for source_row in table.rows:
        target: list[Any] = []
        for index in kept:
            value = source_row[index] if index < len(source_row) else ""
            header = table.headers[index]
            sanitized = sanitize_mapping({header: value})
            target.append(sanitized.get(header, ""))
        if any(value not in (None, "") for value in target):
            rows.append(tuple(target))
    return headers, tuple(rows)


def _render_field_table(document: DocumentType, items: list[tuple[str, Any]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = True
    for index, (key, value) in enumerate(items):
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(cells[0], LIGHT_ACCENT)
        if index % 2:
            _set_cell_shading(cells[1], LIGHT_ROW)
        if cells[0].paragraphs[0].runs:
            cells[0].paragraphs[0].runs[0].bold = True


def _render_data_table(document: DocumentType, source: Table) -> bool:
    headers, rows = _safe_table(source)
    if not headers or not rows:
        return False
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    header_cells = table.rows[0].cells
    for index, label in enumerate(headers):
        header_cells[index].text = str(label)
        _set_cell_shading(header_cells[index], ACCENT)
        if header_cells[index].paragraphs[0].runs:
            header_run = header_cells[index].paragraphs[0].runs[0]
            header_run.bold = True
            header_run.font.color.rgb = RGBColor(255, 255, 255)
    _repeat_table_header(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(row):
            cells[column_index].text = "" if value is None else str(value)
            cells[column_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                _set_cell_shading(cells[column_index], LIGHT_ROW)
    return True


def _render_section(document: DocumentType, section: Section, language: str) -> None:
    document.add_heading(section.title, level=1)
    rendered = False
    items = _safe_field_items(section)
    if items:
        _render_field_table(document, items)
        rendered = True
    for table in section.tables:
        if _render_data_table(document, table):
            rendered = True
    if not rendered:
        paragraph = document.add_paragraph(translate(language, "unavailable"))
        paragraph.runs[0].italic = True


def _render_diagnostics(document: DocumentType, report: ReportData) -> None:
    if not report.diagnostics:
        return
    document.add_heading(translate(report.language, "diagnostics"), level=1)
    table = Table(
        headers=(translate(report.language, "source"), translate(report.language, "message")),
        rows=tuple((item.source, item.message) for item in report.diagnostics),
    )
    _render_data_table(document, table)


def write_docx(report: ReportData, output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.core_properties.title = report.title
    document.core_properties.subject = translate(report.language, "app_name")
    document.core_properties.author = "dsa_cat"
    document.core_properties.comments = "Personal and household use only"
    _configure_styles(document)
    _configure_sections(document, report)

    title = document.add_heading(report.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if report.subtitle:
        identity = document.add_paragraph(style="Report Subtitle")
        identity.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = identity.add_run(report.subtitle)
        run.bold = True
        run.font.size = Pt(13)
    date_line = document.add_paragraph(style="Report Subtitle")
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_line.add_run(f"{translate(report.language, 'report_date')}: {report.generated_at}")
    divider = document.add_paragraph()
    divider.alignment = WD_ALIGN_PARAGRAPH.CENTER
    divider_run = divider.add_run("━━━━━━━━━━━━━━━━━━━━━━━━")
    divider_run.font.color.rgb = RGBColor.from_string(ACCENT)
    toc_label = "Содержание" if report.language == "ru" else "Contents"
    toc_heading = document.add_paragraph()
    toc_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = toc_heading.add_run(toc_label)
    run.bold = True
    run.font.size = Pt(14)
    toc = document.add_paragraph()
    _add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    document.add_page_break()

    for section in report.sections:
        _render_section(document, section, report.language)
    _render_diagnostics(document, report)
    document.save(output)
    return output
